#!/usr/bin/env python3
"""
Rebuild the dashboard data from the combined P&L workbook.

    python scripts/build_data.py path/to/Combined_PnL_by_Store.xlsx

The workbook is the one produced by scripts/merge_pnl.py (one tab per store plus a
Summary tab). The result is written straight into index.html (the <script id="pnl">
block) and also saved as data/pnl.json for reference.
"""
import json, os, re, sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SUBS = {"LABOUR EXPENSE", "OTHER LABOUR EXPENSE"}
VH = {"NET VENDING INCOME"}
AFTER = {"MM EXPENSE": "AUDITOR & SUPPORT COSTS", "UTILITIES": "INSURANCE, SUPPLIES & FEES"}
SPECIAL = {"TOTAL GROSS PROFIT & OTHER INCOME": "GROSS PROFIT & OTHER INCOME",
           "TOTAL ADJUSTED OPERATING EXPENSES": "ADJUSTED OPERATING EXPENSES",
           "NET INCOME": "NET INCOME", "TOTAL FUEL INCOME": "FUEL INCOME"}
ORDER = ["FUEL INCOME", "MERCHANDISE SALES", "MERCHANDISE COST OF SALES", "MERCHANDISE GROSS PROFIT",
         "INVENTORY ADJUSTMENT", "OTHER INCOME", "NET VENDING INCOME", "GROSS PROFIT & OTHER INCOME",
         "LABOUR EXPENSE", "OTHER LABOUR EXPENSE", "MM EXPENSE", "AUDITOR & SUPPORT COSTS", "RENT EXPENSE",
         "REPAIR AND MAINTENANCE", "UTILITIES", "INSURANCE, SUPPLIES & FEES", "OTHER OPERATING EXPENSES",
         "ADJUSTED OPERATING EXPENSES", "NET INCOME"]
META = {  # section -> (role, rows that make up its value)
    "MERCHANDISE SALES": ("income", None), "MERCHANDISE COST OF SALES": ("cogs", None),
    "MERCHANDISE GROSS PROFIT": ("derived", None), "INVENTORY ADJUSTMENT": ("income", None),
    "OTHER INCOME": ("income", None),
    "NET VENDING INCOME": ("income", ["LOTTERY GROSS PROFIT", "LOTTERY VENDING GROSS PROFIT"]),
    "FUEL INCOME": ("income", None), "GROSS PROFIT & OTHER INCOME": ("subtotal", None),
    "LABOUR EXPENSE": ("expense", None), "OTHER LABOUR EXPENSE": ("expense", None),
    "MM EXPENSE": ("expense", None), "AUDITOR & SUPPORT COSTS": ("expense", None),
    "RENT EXPENSE": ("expense", None), "REPAIR AND MAINTENANCE": ("expense", None),
    "UTILITIES": ("expense", None), "INSURANCE, SUPPLIES & FEES": ("expense", None),
    "OTHER OPERATING EXPENSES": ("expense", None),
    "ADJUSTED OPERATING EXPENSES": ("subtotal", None), "NET INCOME": ("subtotal", None)}
LBL = {"MERCHANDISE SALES": "Merchandise Sales", "MERCHANDISE COST OF SALES": "Cost of Sales",
       "MERCHANDISE GROSS PROFIT": "Merchandise Gross Profit", "INVENTORY ADJUSTMENT": "Inventory Adjustment",
       "OTHER INCOME": "Other Income", "NET VENDING INCOME": "Lottery & Vending", "FUEL INCOME": "Fuel Income",
       "GROSS PROFIT & OTHER INCOME": "Gross Profit & Other Income", "LABOUR EXPENSE": "Labour",
       "OTHER LABOUR EXPENSE": "Other Labour", "MM EXPENSE": "Market Manager",
       "AUDITOR & SUPPORT COSTS": "Auditor & Support", "RENT EXPENSE": "Rent",
       "REPAIR AND MAINTENANCE": "Repair & Maintenance", "UTILITIES": "Utilities",
       "INSURANCE, SUPPLIES & FEES": "Insurance, Supplies & Fees", "OTHER OPERATING EXPENSES": "Other Operating",
       "ADJUSTED OPERATING EXPENSES": "Total Operating Expenses", "NET INCOME": "Net Income"}


def build(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    months, stores, lines, line_order = None, [], {}, []
    for tab in wb.sheetnames:
        if tab == "Summary":
            continue
        rows = list(wb[tab].iter_rows(values_only=True))
        full, addr = rows[0][0], rows[1][0]
        ms = [d.strftime("%Y-%m") for d in rows[3][1:8]]
        months = months or ms
        assert ms == months, f"{tab}: months differ from other tabs"
        sid = full.split(" - ")[0].strip()
        short = full.split(" - ", 1)[1].replace(" LLC", "").replace("(MABANK)", "Mabank").title().strip()
        stores.append({"id": sid, "short": short, "name": full, "addr": addr})
        sec, prevk = None, None
        for r in rows[4:]:
            lab = (r[0] or "").strip()
            if not lab:
                continue
            raw = list(r[1:9]); up = lab.upper()
            isnum = [isinstance(v, (int, float)) for v in raw]
            vals = [round(v, 2) if isinstance(v, (int, float)) else 0.0 for v in raw]
            if up in VH or (not any(isnum)):
                sec, prevk = up, "head"
                continue
            if up in SPECIAL:
                k, s = "total", SPECIAL[up]
            elif up.startswith("TOTAL"):
                k, s = "total", sec
            else:
                k = "data"
                s = AFTER.get(sec, "OTHER ITEMS") if prevk == "total" and sec not in AFTER.values() else sec
                if prevk == "total" and sec in AFTER:
                    s = AFTER[sec]
                sec = s
            key = (s, lab)
            if key not in lines:
                lines[key] = {"s": s, "l": lab, "k": k, "by": {}}; line_order.append(key)
            lines[key]["by"][sid] = vals[:7]
            prevk = k
    sids = [s["id"] for s in stores]
    out_lines = [{"s": lines[k]["s"], "l": lines[k]["l"], "k": lines[k]["k"],
                  "v": {i: lines[k]["by"].get(i, [0.0] * 7) for i in sids}} for k in line_order]
    seen = []
    for l in out_lines:
        if l["s"] not in seen:
            seen.append(l["s"])
    unknown = set(seen) - set(ORDER)
    if unknown:
        sys.exit(f"Unknown section(s) in workbook: {sorted(unknown)} - add them to ORDER/META/LBL.")
    secs = [{"k": k, "role": META[k][0], "vr": META[k][1], "n": LBL[k]} for k in ORDER if k in seen]
    out_lines.sort(key=lambda l: ORDER.index(l["s"]))
    return {"months": months, "stores": stores, "sections": secs, "lines": out_lines}


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    data = build(sys.argv[1])
    blob = json.dumps(data, separators=(",", ":")).replace("</", "<\\/")
    os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
    with open(os.path.join(ROOT, "data", "pnl.json"), "w") as f:
        json.dump(data, f, indent=1)
    idx = os.path.join(ROOT, "index.html")
    html = open(idx, encoding="utf-8").read()
    pat = re.compile(r'(<script type="application/json" id="pnl">)(.*?)(</script>)', re.S)
    if not pat.search(html):
        sys.exit("Could not find the data block in index.html")
    html = pat.sub(lambda m: m.group(1) + blob + m.group(3), html, count=1)
    open(idx, "w", encoding="utf-8").write(html)
    print(f"Updated index.html: {len(data['stores'])} stores, {len(data['months'])} months, "
          f"{len(data['lines'])} lines")


if __name__ == "__main__":
    main()
